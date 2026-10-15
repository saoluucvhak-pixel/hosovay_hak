# -*- coding: utf-8 -*-
"""Kiểm tra ĐỘC LẬP đầu ra của kiem_thu/chay_trinh_tao_mau.js: mẫu do trình tạo sinh ra + chứng từ phát hành từ mẫu đó.
Chạy: python3 kiem_thu/kiem_tra_trinh_tao.py <thư_mục_ra>"""
import sys, os, re, json, zipfile, subprocess, warnings
warnings.filterwarnings('ignore')
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
RA = sys.argv[1]; loi = []
def van(doc):
    out = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells: out.extend(p.text for p in c.paragraphs)
    return '\n'.join(out)
def so(s):
    s = re.sub(r'[^\d]', '', s); return int(s) if s else None

# --- mẫu mới ---
m = Document(os.path.join(RA, 'mau_moi.docx')); tm = van(m)
tb = m.tables[1]
if len(tb.rows) != 3: loi.append('mẫu mới: bảng thụ hưởng phải còn 3 hàng (tiêu đề + 1 hàng mẫu + TC), có %d' % len(tb.rows))
for k in ['{{CT.TEN_NGUOI_HUONG}}', '{{CT.SO_TIEN_QUY_DOI}}', '{{CONG_TY.TEN}}', '{{HO_SO.SO_TIEN}}', '{{BANG_KE.TONG_TIEN_QUY_DOI}}']:
    if k not in tm: loi.append('mẫu mới thiếu ' + k)
if 'Chúng tôi cam kết' not in tm or 'không thể tách rời' not in tm: loi.append('mất điều khoản cố định')
if not m.tables[0]._tbl.xpath('.//w:drawing'): loi.append('mất logo')
if 'Hóa đơn GTGT' in tm and '{{HO_SO.TAI_LIEU_CHUNG_MINH}}' not in tm: pass
cel = tb.rows[1].cells
if any(p.alignment != WD_ALIGN_PARAGRAPH.RIGHT for p in cel[3].paragraphs if p.text.strip()): loi.append('ô tiền của mẫu mới không canh phải')
if any(p.alignment != WD_ALIGN_PARAGRAPH.CENTER for p in cel[0].paragraphs if p.text.strip()): loi.append('ô STT của mẫu mới không canh giữa')

# --- chứng từ phát hành từ mẫu mới ---
e = json.load(open(os.path.join(RA, 'chung_tu_moi.json'), encoding='utf8'))
d = Document(os.path.join(RA, 'chung_tu_moi.docx')); td = van(d)
if '{{' in td or '}}' in td: loi.append('chứng từ mới còn {{ }}')
t = d.tables[1]
if len(t.rows) != 1 + e['soDong'] + 1: loi.append('chứng từ mới: số hàng %d' % len(t.rows))
tong = 0
for i in range(e['soDong']):
    c = [x.text for x in t.rows[1 + i].cells]
    if e['ten'][i] not in c[4]: loi.append('dòng %d sai tên: %r' % (i + 1, c[4]))
    tong += so(c[3]) or 0
    if c[0].strip() != '%02d' % (i + 1): loi.append('dòng %d sai STT' % (i + 1))
if tong != e['tong']: loi.append('tổng các dòng %d != %d' % (tong, e['tong']))
tc = [x.text for x in t.rows[-1].cells]
if so(tc[3]) != e['tong']: loi.append('hàng TC sai: %r' % tc)
if e['bangChu'] not in td: loi.append('thiếu tiền bằng chữ')
if '463.520.000' in td: loi.append('lọt số liệu hồ sơ gốc')

# --- PDF qua LibreOffice ---
subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', RA, os.path.join(RA, 'chung_tu_moi.docx')], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=240)
p = os.path.join(RA, 'chung_tu_moi.pdf'); txt = subprocess.run(['pdftotext', '-layout', p, '-'], capture_output=True, text=True).stdout if os.path.exists(p) else ''
if not txt: loi.append('không xuất được PDF')
if '150.000.000' not in txt or 'NGUYỄN VĂN B' not in txt: loi.append('PDF thiếu số liệu hồ sơ mới')
print('KIỂM TRA ĐỘC LẬP:', 'ĐẠT' if not loi else 'LỖI: ' + '; '.join(loi)); print('TỔNG LỖI:', len(loi)); sys.exit(1 if loi else 0)
