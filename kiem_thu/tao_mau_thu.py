# -*- coding: utf-8 -*-
"""Tạo các file mẫu DOCX thử nghiệm (mau_thu/*.docx). Chạy: python3 kiem_thu/tao_mau_thu.py
Mẫu cố ý có: placeholder bị tách thành nhiều run (khác định dạng), header/footer, hàng tiêu đề lặp lại ở mỗi trang,
ô ghép (gridSpan) ở hàng tổng, điều khoản cố định và vùng chữ ký (phải được giữ nguyên)."""
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

RA = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'mau_thu')

DIEU_KHOAN = ('Tôi/Chúng tôi cam kết sẽ bổ sung chứng từ giải ngân trong vòng tối đa 10 ngày làm việc kể từ ngày được giải ngân. '
              'Hợp đồng tín dụng cụ thể này được lập thành 03 bản, là một bộ phận không thể tách rời Hợp đồng tín dụng hạn mức.')


def font(doc):
    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(2)


def runs(p, parts, size=None):
    """parts: list[(text, bold, italic)] — mỗi phần là 1 run riêng (tách placeholder thành nhiều run)."""
    for t, b, i in parts:
        r = p.add_run(t)
        r.bold, r.italic = b, i
        if size:
            r.font.size = Pt(size)
    return p


def header_repeat(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:tblHeader'); e.set(qn('w:val'), 'true'); trPr.append(e)


def row_height(row, cm):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement('w:trHeight'); h.set(qn('w:val'), str(int(cm * 567))); h.set(qn('w:hRule'), 'atLeast'); trPr.append(h)
    c = OxmlElement('w:cantSplit'); trPr.append(c)


def cell(c, parts, align=None, bold=False, size=9):
    p = c.paragraphs[0]
    if isinstance(parts, str):
        parts = [(parts, bold, False)]
    runs(p, parts, size)
    if align:
        p.alignment = align


def set_widths(table, widths):
    table.autofit = False
    for i, w in enumerate(widths):
        table.columns[i].width = Cm(w)          # cập nhật cả tblGrid để Word/LibreOffice dùng đúng độ rộng cột
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)


def tao_de_nghi(duong_dan, usd):
    d = Document(); font(d)
    sec = d.sections[0]
    runs(sec.header.paragraphs[0], [('Số Hợp đồng: {{HOP_DONG', True, False), ('.SO}}', True, False)], 9)      # tách 2 run
    runs(sec.footer.paragraphs[0], [('Mã hồ sơ: ', False, False), ('{{HO_SO', False, True), ('.MA}}', False, True), ('  |  In ngày {{NGAY.HOM_NAY}}', False, False)], 8)

    t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    runs(t, [('HỢP ĐỒNG TÍN DỤNG CỤ THỂ' + (' - PHẦN ĐỀ NGHỊ GIẢI NGÂN' if usd else ''), True, False)], 13)
    runs(d.add_paragraph(), [('Kính gửi: {{HOP_DONG.NGAN_HANG}} - ', False, False), ('{{HOP_DONG.CHI_', True, False), ('NHANH}}', True, False)])
    runs(d.add_paragraph(), [('Bên vay: ', True, False), ('{{CONG_TY.TEN}}', False, False)])
    runs(d.add_paragraph(), [('Địa chỉ: {{CONG_TY.DIA_CHI}}  Điện thoại: {{CONG_TY.DIEN_THOAI}}', False, False)])
    runs(d.add_paragraph(), [('Do ông/bà: {{CONG_TY.NGUOI_DAI_DIEN}}  Chức vụ: {{CONG_TY.CHUC_VU}} làm đại diện', False, False)])
    runs(d.add_paragraph(), [('Căn cứ: Hợp đồng tín dụng hạn mức số {{HOP_DONG.SO}} ký ngày {{HOP_DONG.NGAY}}', False, True)])
    runs(d.add_paragraph(), [('Đề nghị rút số tiền vay là: ', False, False), ('{{HO_SO.SO_', True, False), ('TIEN}}', True, False),
                             (' {{HO_SO.DON_VI_TIEN}} (Bằng chữ: ', False, False), ('{{HO_SO.SO_TIEN_BANG', False, True), ('_CHU}}', False, True), (')', False, False)])
    if usd:
        runs(d.add_paragraph(), [('Tỷ giá bán: {{HO_SO.TY_GIA}} VND/USD. Số tiền quy đổi: {{HO_SO.SO_TIEN_QUY_DOI}} đồng ({{HO_SO.SO_TIEN_QUY_DOI_BANG_CHU}}).', False, False)])
    runs(d.add_paragraph(), [('Thời hạn vay: {{HO_SO.THOI_HAN_NGAY}} ngày. Ngày giải ngân: {{HO_SO.NGAY_GIAI_NGAN}}. Đến hạn: {{HO_SO.NGAY_DEN_HAN}}.', False, False)])
    runs(d.add_paragraph(), [('Để thanh toán theo nội dung dưới đây:', False, False)])

    cot = 7 if usd else 6
    tb = d.add_table(rows=3, cols=cot); tb.style = 'Table Grid'; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    tieu_de = ['TT', 'Nội dung', 'Loại chứng từ', 'Số hiệu chứng từ', 'Số tiền' + (' (USD)' if usd else ' (đồng)')]
    if usd:
        tieu_de.append('Quy đổi (đồng)')
    tieu_de.append('Đơn vị thụ hưởng')
    for i, h in enumerate(tieu_de):
        cell(tb.rows[0].cells[i], h, WD_ALIGN_PARAGRAPH.CENTER, True)
    header_repeat(tb.rows[0])
    # hàng mẫu (placeholder tách run ở vài ô)
    r = tb.rows[1]; row_height(r, 1.3)
    cell(r.cells[0], '{{CT.STT_2}}', WD_ALIGN_PARAGRAPH.CENTER)
    cell(r.cells[1], [('{{CT.NOI_', False, False), ('DUNG}}', False, False)], WD_ALIGN_PARAGRAPH.LEFT)
    cell(r.cells[2], '{{CT.LOAI_CHUNG_TU}}', WD_ALIGN_PARAGRAPH.LEFT)
    cell(r.cells[3], '{{CT.SO_TAI_LIEU}}', WD_ALIGN_PARAGRAPH.LEFT)
    cell(r.cells[4], [('{{CT.SO', True, False), ('_TIEN}}', True, False)], WD_ALIGN_PARAGRAPH.RIGHT)
    k = 5
    if usd:
        cell(r.cells[5], '{{CT.SO_TIEN_QUY_DOI}}', WD_ALIGN_PARAGRAPH.RIGHT); k = 6
    cell(r.cells[k], '{{CT.TEN}}\nSố TK: {{CT.SO_TK}}\nTại: {{CT.NGAN_HANG}}', WD_ALIGN_PARAGRAPH.LEFT)
    # hàng tổng: ghép 4 ô đầu
    tong = tb.rows[2]
    m = tong.cells[0].merge(tong.cells[3])
    cell(m, 'TỔNG CỘNG', WD_ALIGN_PARAGRAPH.CENTER, True)
    cell(tong.cells[4], '{{BANG_KE.TONG_TIEN}}', WD_ALIGN_PARAGRAPH.RIGHT, True)
    if usd:
        cell(tong.cells[5], '{{BANG_KE.TONG_TIEN_QUY_DOI}}', WD_ALIGN_PARAGRAPH.RIGHT, True)
    set_widths(tb, [1.0, 3.6, 2.4, 2.6, 2.4, 2.4, 4.0] if usd else [1.0, 4.0, 2.6, 2.8, 2.6, 4.4])

    d.add_paragraph()
    runs(d.add_paragraph(), [('Tài liệu liên quan: {{HO_SO.TAI_LIEU_CHUNG_MINH}}', False, False)])
    runs(d.add_paragraph(), [('Chúng tôi cam kết:', True, True)])
    runs(d.add_paragraph(), [(DIEU_KHOAN, False, False)])
    if usd:
        runs(d.add_paragraph(), [('Tôi/Chúng tôi đồng ý bán giao ngay toàn bộ số tiền USD được vay cho Ngân hàng theo tỷ giá bán: {{HO_SO.TY_GIA}} VND/USD.', False, False)])
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    runs(p, [('Đà Nẵng, {{NGAY.GIAI_NGAN_DAY_DU}}', False, True)])
    ky = d.add_table(rows=1, cols=2); ky.style = 'Table Grid'
    cell(ky.rows[0].cells[0], 'KẾ TOÁN TRƯỞNG\n(Ký, họ tên)', WD_ALIGN_PARAGRAPH.CENTER, True, 11)
    cell(ky.rows[0].cells[1], 'BÊN VAY\n(ký, ghi rõ họ tên và đóng dấu)\n\n\n\n{{CONG_TY.NGUOI_DAI_DIEN}}', WD_ALIGN_PARAGRAPH.CENTER, True, 11)
    d.save(duong_dan)


def tao_unc_mot_khoan(duong_dan, usd):
    """Ủy nhiệm chi 1 khoản — placeholder {{CT.*}} nằm NGOÀI bảng (áp dụng khi hồ sơ có đúng 1 dòng)."""
    d = Document(); font(d)
    runs(d.sections[0].footer.paragraphs[0], [('UNC - {{HO_SO.MA}}', False, False)], 8)
    t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    runs(t, [('ỦY NHIỆM CHI / PAYMENT ORDER', True, False)], 14)
    runs(d.add_paragraph(), [('Ngày/Date: {{HO_SO.NGAY_GIAI_NGAN}}', False, False)])
    runs(d.add_paragraph(), [('Tên tài khoản trích nợ: {{CONG_TY.TEN}}', False, False)])
    runs(d.add_paragraph(), [('Số tài khoản trích nợ: ' + ('{{CONG_TY.TK_USD}}' if usd else '{{CONG_TY.TK_VND}}'), False, False)])
    runs(d.add_paragraph(), [('Số tiền bằng số: ', False, False), ('{{CT.SO_TIEN}} ' + ('USD' if usd else 'đồng'), True, False)])
    runs(d.add_paragraph(), [('Số tiền bằng chữ: ', False, False), ('{{CT.SO_TIEN_BANG_', False, True), ('CHU}}', False, True)])
    runs(d.add_paragraph(), [('Người hưởng: {{CT.TEN}}', False, False)])
    runs(d.add_paragraph(), [('Số tài khoản: {{CT.SO_TK}}  Tại ngân hàng: {{CT.NGAN_HANG}}', False, False)])
    runs(d.add_paragraph(), [('Nội dung: {{CT.NOI_DUNG}}', False, False)])
    ky = d.add_table(rows=1, cols=2); ky.style = 'Table Grid'
    cell(ky.rows[0].cells[0], 'Kế toán trưởng', WD_ALIGN_PARAGRAPH.CENTER, True, 11)
    cell(ky.rows[0].cells[1], 'Chủ tài khoản', WD_ALIGN_PARAGRAPH.CENTER, True, 11)
    d.save(duong_dan)


def tao_bang_ke(duong_dan, usd):
    d = Document(); font(d)
    sec = d.sections[0]
    runs(sec.footer.paragraphs[0], [('Bảng kê - {{HO_SO.MA}}', False, False)], 8)
    runs(d.add_paragraph(), [('Tên đơn vị: {{CONG_TY.TEN}}', False, False)])
    t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    runs(t, [('BẢNG KÊ CHUYỂN TIỀN VAY NGÀY {{HO_SO.NGAY_GIAI_NGAN}}', True, False)], 13)
    cot = 6 if usd else 5
    tb = d.add_table(rows=3, cols=cot); tb.style = 'Table Grid'
    th = ['Stt', 'Tên đơn vị thụ hưởng', 'Số tài khoản', 'Ngân hàng', 'Số tiền' + (' (USD)' if usd else ' (VND)')]
    if usd:
        th.append('Quy đổi (VND)')
    for i, h in enumerate(th):
        cell(tb.rows[0].cells[i], h, WD_ALIGN_PARAGRAPH.CENTER, True)
    header_repeat(tb.rows[0])
    r = tb.rows[1]; row_height(r, 1.1)
    cell(r.cells[0], '{{CT.STT}}', WD_ALIGN_PARAGRAPH.CENTER)
    cell(r.cells[1], '{{CT.TEN}}', WD_ALIGN_PARAGRAPH.LEFT)
    cell(r.cells[2], '{{CT.SO_TK}}', WD_ALIGN_PARAGRAPH.LEFT)
    cell(r.cells[3], '{{CT.NGAN_HANG}}', WD_ALIGN_PARAGRAPH.LEFT)
    cell(r.cells[4], '{{CT.SO_TIEN}}', WD_ALIGN_PARAGRAPH.RIGHT)
    if usd:
        cell(r.cells[5], '{{CT.SO_TIEN_QUY_DOI}}', WD_ALIGN_PARAGRAPH.RIGHT)
    tong = tb.rows[2]
    m = tong.cells[0].merge(tong.cells[3])
    cell(m, 'Tổng cộng', WD_ALIGN_PARAGRAPH.CENTER, True)
    cell(tong.cells[4], '{{BANG_KE.TONG_TIEN}}', WD_ALIGN_PARAGRAPH.RIGHT, True)
    if usd:
        cell(tong.cells[5], '{{BANG_KE.TONG_TIEN_QUY_DOI}}', WD_ALIGN_PARAGRAPH.RIGHT, True)
    runs(d.add_paragraph(), [('Bằng chữ: {{BANG_KE.TONG_TIEN_BANG_CHU}}', False, True)])
    d.save(duong_dan)


def tao_hdmb(duong_dan):
    d = Document(); font(d)
    runs(d.sections[0].footer.paragraphs[0], [('HDMB - {{HO_SO.MA}}', False, False)], 8)
    t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    runs(t, [('HỢP ĐỒNG MUA BÁN NGOẠI TỆ GIAO NGAY', True, False)], 14)
    runs(d.add_paragraph(), [('Căn cứ Hợp đồng khung số {{HOP_DONG.HD_KHUNG_MBNT_SO}} ngày {{HOP_DONG.HD_KHUNG_MBNT_NGAY}} giữa {{HOP_DONG.NGAN_HANG}} – {{HOP_DONG.CHI_NHANH}} và {{CONG_TY.TEN}}.', False, True)])
    runs(d.add_paragraph(), [('Số tiền giao dịch: ', False, False), ('{{HO_SO.SO_TIEN}} USD', True, False), (' (Bằng chữ: {{HO_SO.SO_TIEN_BANG_CHU}})', False, False)])
    runs(d.add_paragraph(), [('Tỷ giá: {{HO_SO.TY_GIA}} VND/USD. Số tiền thanh toán: {{HO_SO.SO_TIEN_QUY_DOI}} đồng', False, False)])
    runs(d.add_paragraph(), [('Bằng chữ: {{HO_SO.SO_TIEN_QUY_DOI_BANG_CHU}}', False, True)])
    runs(d.add_paragraph(), [('Ngày giao dịch: {{HO_SO.NGAY_GIAI_NGAN}}. Tài khoản USD: {{CONG_TY.TK_USD}}. Tài khoản VND: {{CONG_TY.TK_VND}}.', False, False)])
    d.save(duong_dan)


if __name__ == '__main__':
    os.makedirs(RA, exist_ok=True)
    tao_de_nghi(os.path.join(RA, 'DE_NGHI_USD.docx'), True)
    tao_de_nghi(os.path.join(RA, 'DE_NGHI_VND.docx'), False)
    tao_unc_mot_khoan(os.path.join(RA, 'UNC_USD_1_KHOAN.docx'), True)
    tao_unc_mot_khoan(os.path.join(RA, 'UNC_VND_1_KHOAN.docx'), False)
    tao_bang_ke(os.path.join(RA, 'BANG_KE_USD.docx'), True)
    tao_bang_ke(os.path.join(RA, 'BANG_KE_VND.docx'), False)
    tao_hdmb(os.path.join(RA, 'HDMB_USD.docx'))
    print('Đã tạo mẫu thử tại', os.path.abspath(RA))
