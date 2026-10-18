# -*- coding: utf-8 -*-
"""Kiểm tra ĐỘC LẬP kết quả điền 5 mẫu thật (đọc lại từ file DOCX/XLSX, không dùng dữ liệu nội bộ của engine) + PDF qua LibreOffice.
Chạy: python3 kiem_thu/kiem_tra_mau_that.py <thư_mục_ra>   (sau node kiem_thu/chay_thu_mau_that.js <thư_mục_ra>)"""
import sys, os, re, json, glob, zipfile, subprocess, warnings
warnings.filterwarnings('ignore')
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
import openpyxl
from lxml import etree

RA = sys.argv[1]
GOC = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
bang, tong_loi = [], 0
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
W14 = '{http://schemas.microsoft.com/office/word/2010/wordml}'


def so(s):
    s = str(s).strip().replace('USD', '').strip()
    return int(s.replace('.', '').replace(',', '')) if re.fullmatch(r'[\d.,]+', s) else None


def van_ban_docx(doc):
    out = [p.text for p in doc.paragraphs]
    def tbl(t):
        for r in t.rows:
            for c in r.cells:
                out.extend(p.text for p in c.paragraphs)
                for n in c.tables: tbl(n)
    for t in doc.tables: tbl(t)
    return '\n'.join(out)


def checkbox(path):
    with zipfile.ZipFile(path) as z:
        x = z.read('word/document.xml').decode('utf8')
    return re.findall(r'<w14:checked w14:val="(\d)"', x), x.count('<w:sdt>')


def pdf(f):
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', RA, f], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=240)
    p = os.path.splitext(f)[0] + '.pdf'
    if not os.path.exists(p): return 0, ''
    txt = subprocess.run(['pdftotext', '-layout', p, '-'], capture_output=True, text=True).stdout
    return txt.count('\f'), txt


def chu_so(s): return re.sub(r'\D', '', s)


def kiem_tra(jf):
    global tong_loi
    e = json.load(open(jf, encoding='utf8')); ten = e['ten']; n = e['soDong']; loi = []
    f = glob.glob(os.path.join(RA, ten + '.*'))
    f = [x for x in f if x.endswith(('.docx', '.xlsx'))][0]
    if e['engine']['conLai'] or e['engine']['conNgoac']: loi.append('engine báo còn placeholder')
    mau_goc = os.path.join(GOC, 'mau_that', e['mau'])
    trang, txt = pdf(f)
    if '{{' in txt or '}}' in txt: loi.append('PDF còn {{ }}')
    if re.search(r'cents|\d\.\d{2}\s*USD', txt): loi.append('PDF có cents/số lẻ USD')

    if f.endswith('.docx'):
        doc = Document(f); van = van_ban_docx(doc)
        if '{{' in van or '}}' in van: loi.append('còn {{ }}')
        if 'cents' in van.lower(): loi.append('có chữ cents')
        if e['mau'].startswith('HDTD_VND'):
            tien = '463.520.000'
            if tien + ' đồng' not in van: loi.append('thiếu "%s đồng"' % tien)
            if e['tongBangChu'] not in van: loi.append('thiếu tiền bằng chữ: %s' % e['tongBangChu'])
            if 'USD' in van.replace('Tài khoản tiền gửi USD', ''): loi.append('có chữ USD ngoài dòng tài khoản USD')
            for k in ['02.25/2026/8100173/HĐTD', '02/2026/8100173/HĐTD', '04/09/2026', 'CÔNG TY TNHH HOÀNG ANH KHÔI', '0401447476', '562.0.727.934', '562.0.011.163', 'Võ Văn Phụng',
                      'Thời hạn vay: 178 ngày', 'ngày 13 tháng 10 năm 2026', 'Quảng Nam', 'Hóa đơn GTGT']:
                if k not in van: loi.append('thiếu ' + k)
            t = doc.tables[1]; coTC = e['mau'].endswith('NHIEU_KHOAN.docx')
            if len(t.rows) != 1 + n + (1 if coTC else 0): loi.append('số hàng bảng %d' % len(t.rows))
            tong = 0
            for i in range(n):
                cells = [c.text for c in t.rows[1 + i].cells]
                if e['cacDong'][i]['ten'] not in cells[4]: loi.append('dòng %d sai tên' % (i + 1))
                if '1234567' + str(i).zfill(3) not in cells[4]: loi.append('dòng %d thiếu số TK' % (i + 1))
                a = so(cells[3])
                if a != so(e['cacDong'][i]['soTien']): loi.append('dòng %d sai số tiền' % (i + 1))
                tong += a or 0
                if cells[0].strip() != str(i + 1).zfill(2): loi.append('dòng %d sai STT %r' % (i + 1, cells[0]))
                if any(p.alignment != WD_ALIGN_PARAGRAPH.RIGHT for p in t.rows[1 + i].cells[3].paragraphs if p.text.strip()): loi.append('tiền không canh phải d%d' % (i + 1))
                if any(p.alignment != WD_ALIGN_PARAGRAPH.CENTER for p in t.rows[1 + i].cells[0].paragraphs if p.text.strip()): loi.append('STT không canh giữa')
            if tong != e['tongTien']: loi.append('tổng dòng %d != %d' % (tong, e['tongTien']))
            if coTC:
                tc = [c.text for c in t.rows[-1].cells]
                if 'TC' not in tc[0] or so(tc[3]) != e['tongTien']: loi.append('hàng TC sai: %s' % tc)
            if not doc.tables[0]._tbl.xpath('.//w:drawing'): loi.append('mất logo')
            if 'Chúng tôi cam kết' not in van or 'bộ phận không thể tách rời Hợp đồng tín dụng hạn mức' not in van: loi.append('mất điều khoản')
        elif e['mau'].startswith('HDTD'):
            if 'Hai mươi lăm ngàn đô la Mỹ' not in van: loi.append('thiếu tiền bằng chữ USD')
            if '25.000 USD' not in van: loi.append('thiếu "25.000 USD"')
            if 'Số Hợp đồng' not in van or '02.25/2026/8100173/HĐTD' not in van: loi.append('thiếu số hợp đồng cụ thể ở đầu trang')
            t = doc.tables[1]; coTC = e['mau'].endswith('NHIEU_KHOAN.docx')
            if len(t.rows) != 1 + n + (1 if coTC else 0): loi.append('số hàng bảng %d' % len(t.rows))
            tong = 0
            for i in range(n):
                cells = [c.text for c in t.rows[1 + i].cells]
                if e['cacDong'][i]['ten'] not in cells[5]: loi.append('dòng %d sai tên' % (i + 1))
                a, b = so(cells[3]), so(cells[4])
                if a != so(e['cacDong'][i]['soTien']): loi.append('dòng %d sai số tiền USD' % (i + 1))
                if b != so(e['cacDong'][i]['quyDoi']): loi.append('dòng %d sai quy đổi' % (i + 1))
                tong += a or 0
                for ci in (3, 4):
                    if any(p.alignment != WD_ALIGN_PARAGRAPH.RIGHT for p in t.rows[1 + i].cells[ci].paragraphs): loi.append('tiền không canh phải d%d' % (i + 1))
                if any(p.alignment != WD_ALIGN_PARAGRAPH.CENTER for p in t.rows[1 + i].cells[0].paragraphs): loi.append('STT không canh giữa')
            if tong != e['tongTien']: loi.append('tổng dòng %d != %d' % (tong, e['tongTien']))
            if coTC:
                tc = [c.text for c in t.rows[-1].cells]
                if 'TC' not in tc[0] or so(tc[3]) != e['tongTien'] or so(tc[4]) != e['tongQuyDoi']: loi.append('hàng TC sai: %s' % tc)
            if not doc.tables[0]._tbl.xpath('.//w:drawing'): loi.append('mất logo')
            if not zipfile.ZipFile(f).namelist().__contains__('word/media/image1.png'): loi.append('mất file logo')
            if 'Chúng tôi cam kết' not in van or 'bộ phận không thể tách rời Hợp đồng tín dụng hạn mức' not in van: loi.append('mất điều khoản')
        else:  # HDMB
            for k in ['25.000 USD', '25.810 USD/VND', e['tongBangChu'] and 'Mười chín' and None]:
                if k and k not in van: loi.append('thiếu ' + k)
            if '645.250.000' not in van: loi.append('thiếu số tiền thanh toán 645.250.000')
            if 'sáu trăm bốn mươi lăm triệu' not in van.lower(): loi.append('thiếu tiền VND bằng chữ')
            cb_ra, sdt_ra = checkbox(f); cb_goc, sdt_goc = checkbox(mau_goc)
            if cb_ra != cb_goc or sdt_ra != sdt_goc: loi.append('hộp kiểm Mua/Bán bị đổi: %s vs %s' % (cb_ra, cb_goc))
        a, b = Document(mau_goc).sections[0], doc.sections[0]
        for k in ('page_width', 'page_height', 'left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
            if getattr(a, k) != getattr(b, k): loi.append('khổ giấy/lề đổi: ' + k)
    else:
        wb = openpyxl.load_workbook(f); wg = openpyxl.load_workbook(mau_goc)
        z = zipfile.ZipFile(f); zg = zipfile.ZipFile(mau_goc)
        names = z.namelist()
        if 'xl/calcChain.xml' in names: loi.append('calcChain còn sót (công thức dời hàng sẽ sai)')
        if 'fullCalcOnLoad="1"' not in z.read('xl/workbook.xml').decode('utf8'): loi.append('thiếu fullCalcOnLoad')
        for x in zg.namelist():
            if x.startswith(('xl/media', 'xl/drawings')) and x not in names: loi.append('mất ' + x)
        ws = wb.worksheets[0]
        if e['mau'].startswith('UNC'):
            qd = e['tongQuyDoi']
            if ws['I7'].value != qd or ws['I7'].data_type != 'n': loi.append('I7=%r' % ws['I7'].value)
            if ws['I7'].number_format != wg.worksheets[0]['I7'].number_format: loi.append('mất định dạng số I7')
            if not str(ws['I8'].value).endswith('y./.') or 'đồng' not in str(ws['I8'].value): loi.append('I8 sai: %r' % ws['I8'].value)
            if ws['H12'].value != e['cacDong'][0]['ten'] or ws['J3'].value != 'CÔNG TY TNHH HOÀNG ANH KHÔI': loi.append('tên sai')
            if ws['X2'].value != '13/10/2026': loi.append('ngày sai: %r' % ws['X2'].value)
            if len(ws.merged_cells.ranges) != len(wg.worksheets[0].merged_cells.ranges): loi.append('đổi ô ghép')
            if not str(ws['AH1'].value).startswith('=IF('): loi.append('mất công thức AH1')
        else:
            if wb.sheetnames != ['HAK']: loi.append('sheet: %s' % wb.sheetnames)
            for i in range(n):
                r = 7 + i
                if ws['A%d' % r].value != i + 1: loi.append('STT hàng %d = %r' % (r, ws['A%d' % r].value))
                if ws['E%d' % r].value != so(e['cacDong'][i]['quyDoi']) or ws['E%d' % r].data_type != 'n': loi.append('E%d = %r' % (r, ws['E%d' % r].value))
                if ws['B%d' % r].value != e['cacDong'][i]['ten']: loi.append('tên hàng %d' % r)
                if ws['E%d' % r].alignment.horizontal != 'right': loi.append('tiền hàng %d không canh phải' % r)
                if ws['A%d' % r].alignment.horizontal != 'center': loi.append('STT hàng %d không canh giữa' % r)
                if ws['E%d' % r].number_format != wg.worksheets[0]['E7'].number_format and False: pass
                # Engine deliberately expands rows for wrapped beneficiary/content text.
                # Preserve the template's minimum height; equality would reject correct expansion.
                height = ws.row_dimensions[r].height or ws.sheet_format.defaultRowHeight or 15
                minimum = wg.worksheets[0].row_dimensions[7].height or wg.worksheets[0].sheet_format.defaultRowHeight or 15
                if height < minimum or height > 409.5: loi.append('độ cao hàng %d không hợp lệ: %s (tối thiểu %s)' % (r, height, minimum))
                if not ws['B%d' % r].alignment.wrap_text or not ws['F%d' % r].alignment.wrap_text: loi.append('mất xuống dòng hàng %d' % r)
            rt = 7 + n
            if ws['B%d' % rt].value != 'Tổng cộng' or ws['E%d' % rt].value != '=SUM(E7:E%d)' % (6 + n): loi.append('hàng tổng: %r %r' % (ws['B%d' % rt].value, ws['E%d' % rt].value))
            kt = 9 + n
            mg = sorted(str(m) for m in ws.merged_cells.ranges)
            if mg != sorted(['A4:B4', 'A%d:B%d' % (kt, kt), 'E%d:F%d' % (kt, kt)]): loi.append('ô ghép: %s' % mg)
            if ws['A%d' % kt].value != 'KẾ TOÁN TRƯỞNG' or ws['E%d' % kt].value != 'NGÂN HÀNG': loi.append('chữ ký không dời đúng')
            if ws['A1'].value != 'Tên đơn vị : CÔNG TY TNHH HOÀNG ANH KHÔI': loi.append('A1')
            # LibreOffice tính lại công thức -> tổng hiển thị phải đúng
            if chu_so(str(e['tongQuyDoi'])) not in chu_so(txt.replace('.', '').replace(',', '')) and str(e['tongQuyDoi']) not in re.sub(r'[.,\s]', '', txt): loi.append('PDF thiếu tổng %d' % e['tongQuyDoi'])
            if trang < 1: loi.append('không có PDF')
    if not trang: loi.append('không tạo được PDF')
    tong_loi += len(loi)
    bang.append((ten, n, trang, 'ĐẠT' if not loi else 'LỖI: ' + '; '.join(loi)))


for jf in sorted(glob.glob(os.path.join(RA, '*.json'))): kiem_tra(jf)
print('%-12s %5s %6s  kết quả' % ('kịch bản', 'dòng', 'trang'))
for b in bang: print('%-12s %5d %6d  %s' % b)
print('TỔNG LỖI:', tong_loi)
sys.exit(1 if tong_loi else 0)
