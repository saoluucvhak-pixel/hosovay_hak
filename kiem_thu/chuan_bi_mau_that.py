# -*- coding: utf-8 -*-
"""Chuẩn bị MẪU THẬT (có placeholder) từ 4 file biểu mẫu BIDV của HAK: đặt {{...}} thay cho dữ liệu mẫu, GIỮ NGUYÊN mọi định dạng khác.

  python3 kiem_thu/chuan_bi_mau_that.py <thư_mục_chứa_4_file_gốc> [thư_mục_ra=mau_that]

File gốc (tự dò theo tên): HDTD_USD_cu_the_*.docx · HD_MBNT_GIAO_NGAY_*.docx · UNC.xlsx · BANG_KE_GIAI_NGAN*.xlsx
Kết quả: HDTD_USD_1_KHOAN.docx, HDTD_USD_NHIEU_KHOAN.docx, HDMB_USD.docx, UNC_1_KHOAN.xlsx, BANG_KE_NHIEU_KHOAN.xlsx
Quy ước theo yêu cầu: tiền canh phải, STT/ngày canh giữa, chuỗi canh trái (áp dụng cho các ô số liệu của bảng thụ hưởng).
Không đưa file gốc vào repository — chỉ các mẫu đã đặt placeholder.
"""
import sys, os, re, glob, copy, zipfile, shutil, tempfile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from lxml import etree
import openpyxl
from openpyxl.styles import Alignment

SRC = sys.argv[1]
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'mau_that')
os.makedirs(OUT, exist_ok=True)


def tim(mau):
    f = glob.glob(os.path.join(SRC, mau))
    assert f, 'không thấy file ' + mau
    return f[0]


# ---------------------------------------------------------------- DOCX: thay chuỗi xuyên nhiều run
def thay(p, cu, moi, regex=False):
    """Thay `cu` (chuỗi hoặc regex) trong đoạn p bằng `moi`; chữ mới nằm ở run đầu tiên chạm tới (giữ định dạng run đó)."""
    runs = p.runs
    texts = [r.text for r in runs]
    full = ''.join(texts)
    if regex:
        m = re.search(cu, full)
        assert m, 'không thấy mẫu regex %r trong %r' % (cu, full)
        i, n = m.start(), m.end() - m.start()
    else:
        i = full.find(cu)
        assert i >= 0, 'không thấy %r trong %r' % (cu, full)
        n = len(cu)
    pos, dau = 0, None
    for k, r in enumerate(runs):
        s, e = pos, pos + len(texts[k]); pos = e
        a, b = max(s, i), min(e, i + n)
        if a < b:
            r.text = texts[k][:a - s] + (moi if dau is None else '') + texts[k][b - s:]
            if dau is None:
                dau = k


def doan_chua(container, chuoi):
    for p in container.paragraphs:
        if chuoi in p.text:
            return p
    raise AssertionError('không thấy đoạn chứa %r' % chuoi)


def canh(p, kieu):
    p.alignment = {'phai': WD_ALIGN_PARAGRAPH.RIGHT, 'giua': WD_ALIGN_PARAGRAPH.CENTER, 'trai': WD_ALIGN_PARAGRAPH.LEFT}[kieu]


def dat_noi_dung_o(cell, text):
    """Đặt chữ cho 1 ô: giữ đoạn/run đầu có chữ, xoá chữ ở các run/đoạn khác (giữ định dạng ô)."""
    dau = False
    for p in cell.paragraphs:
        for r in p.runs:
            if not dau and r.text.strip() != '':
                r.text = text; dau = True
            else:
                r.text = ''
    if not dau:
        p = cell.paragraphs[-1]
        (p.runs[0] if p.runs else p.add_run()).text = text


# ---------------------------------------------------------------- HDTD (Hợp đồng tín dụng cụ thể – phần đề nghị giải ngân, USD)
def hdtd():
    d = Document(tim('*HDTD_USD_cu_the*.docx'))
    # tiêu đề: "Số Hợp đồng" của hợp đồng tín dụng CỤ THỂ = số giấy nhận nợ
    for c in d.tables[0].rows[0].cells:
        for p in c.paragraphs:
            if 'Số Hợp đồng' in p.text:
                thay(p, r'02\.25/2026/8100173/HĐTD', '{{HO_SO.SO_GIAY_NHAN_NO}}', True)
    P = {i: p for i, p in enumerate(d.paragraphs)}
    thay(doan_chua(d, 'Kính gửi:'), 'Ngân hàng TMCP Đầu tư và Phát triển Việt Nam - Chi nhánh Quảng Nam', '{{HOP_DONG.TEN_NGAN_HANG}} - {{HOP_DONG.CHI_NHANH}}')
    thay(doan_chua(d, 'Bên vay:'), 'Công ty TNHH Hoàng Anh Khôi', '{{CONG_TY.TEN}}')
    thay(doan_chua(d, 'Địa chỉ:'), r'Thửa đất số 1182.*Việt Nam', '{{CONG_TY.DIA_CHI}}', True)
    thay(doan_chua(d, 'Điện thoại:'), '0913.413.214', '{{CONG_TY.DIEN_THOAI}}')
    thay(doan_chua(d, 'Giấy chứng nhận đăng ký doanh nghiệp'), '0401447476', '{{CONG_TY.MA_SO_THUE}}')
    thay(doan_chua(d, 'Tài khoản tiền gửi VND'), '562.0.727.934', '{{CONG_TY.TAI_KHOAN_VND}}')
    thay(doan_chua(d, 'Tài khoản tiền gửi USD'), '562.0.011.163', '{{CONG_TY.TAI_KHOAN_USD}}')
    p = doan_chua(d, 'Do ông:')
    thay(p, 'Võ Văn Phụng', '{{CONG_TY.NGUOI_DAI_DIEN}}'); thay(p, 'Giám đốc', '{{CONG_TY.CHUC_VU}}')
    p = doan_chua(d, 'Căn cứ: Hợp đồng tín dụng hạn mức')
    thay(p, '02/2026/8100173/HĐTD', '{{HOP_DONG.SO_HOP_DONG}}'); thay(p, '04/09/2026', '{{HOP_DONG.NGAY_KY}}')
    p = doan_chua(d, 'Đề nghị Ngân hàng cho')
    thay(p, r'19,372\.34 USD', '{{HO_SO.SO_TIEN}} {{HO_SO.DON_VI_TIEN}}', True)
    thay(p, r'Mười chín ngàn.*cents\./\.', '{{HO_SO.SO_TIEN_BANG_CHU}}./.', True)
    thay(doan_chua(d, 'Thời hạn vay'), r'178 ngày', '{{HO_SO.THOI_HAN_VAY}} ngày', True)
    thay(doan_chua(d, 'Ngày vay'), '13/10/2026', '{{HO_SO.NGAY_GIAI_NGAN}}')
    thay(doan_chua(d, 'Tài liệu liên quan'), 'Tài liệu liên quan: ', 'Tài liệu liên quan: {{HO_SO.TAI_LIEU_CHUNG_MINH}}')
    p = doan_chua(d, 'bán giao ngay toàn bộ số tiền USD')
    thay(p, r'01…\./202\.\./8100173/HĐTD', '{{HO_SO.SO_GIAY_NHAN_NO}}', True)
    thay(p, r'tỷ giá …+ USD/VND', 'tỷ giá {{HO_SO.TY_GIA}} USD/VND', True)
    thay(p, r'tại ngày …+\.', 'tại ngày {{HO_SO.NGAY_GIAI_NGAN}}', True)
    thay(doan_chua(d, 'tháng .......'), r', ngày \.+ tháng \.+ năm …', ', {{NGAY.GIAI_NGAN_DAY_DU}}', True)
    # bảng thụ hưởng: 1 hàng mẫu (hàng 1)
    t = d.tables[1]; r = t.rows[1].cells
    thay(r[0].paragraphs[1], '01', '{{CT.STT_2}}')
    thay(r[1].paragraphs[0], 'Giải ngân mua keo lá tràm', '{{CT.NOI_DUNG}}')
    thay(r[2].paragraphs[1], r'UNC ngày 13/10/2026', '{{CT.TAI_LIEU_NGAY}}', True)
    thay(r[3].paragraphs[1], r'19,372\.34 USD', '{{CT.SO_TIEN}} USD', True)
    thay(r[4].paragraphs[1], r'500,000,095', '{{CT.SO_TIEN_QUY_DOI}}', True)
    thay(r[5].paragraphs[0], 'Công Ty TNHH Hoàng Anh Khôi', '{{CT.TEN_NGUOI_HUONG}}')
    thay(r[5].paragraphs[1], '8619299999', '{{CT.SO_TAI_KHOAN}}')
    thay(r[5].paragraphs[2], 'BIDV _ CN Quảng Nam', '{{CT.NGAN_HANG}}')
    # canh lề theo quy ước: STT giữa · tiền phải · chuỗi trái
    for p in r[0].paragraphs: canh(p, 'giua')
    for ci in (3, 4):
        for p in r[ci].paragraphs: canh(p, 'phai')
    for p in r[1].paragraphs + r[5].paragraphs + r[2].paragraphs: canh(p, 'trai')
    d.save(os.path.join(OUT, 'HDTD_USD_1_KHOAN.docx'))

    # biến thể NHIỀU KHOẢN: thêm hàng "TC" (tổng cộng) — nhân bản hàng mẫu rồi thay nội dung
    d2 = Document(os.path.join(OUT, 'HDTD_USD_1_KHOAN.docx'))
    t2 = d2.tables[1]
    tc = copy.deepcopy(t2.rows[1]._tr)
    t2.rows[1]._tr.addnext(tc)
    from docx.table import _Row
    hang_tc = _Row(tc, t2)
    c = hang_tc.cells
    dat_noi_dung_o(c[0], 'TC')
    for p in c[0].paragraphs:
        for rr in p.runs: rr.bold = True
    dat_noi_dung_o(c[1], ''); dat_noi_dung_o(c[2], '')
    dat_noi_dung_o(c[3], '{{BANG_KE.TONG_TIEN}} USD'); dat_noi_dung_o(c[4], '{{BANG_KE.TONG_TIEN_QUY_DOI}}'); dat_noi_dung_o(c[5], '')
    for ci in (3, 4):
        for p in c[ci].paragraphs:
            canh(p, 'phai')
            for rr in p.runs: rr.bold = True
    d2.save(os.path.join(OUT, 'HDTD_USD_NHIEU_KHOAN.docx'))


# ---------------------------------------------------------------- HDMB (Hợp đồng mua bán ngoại tệ giao ngay)
def hdmb():
    d = Document(tim('*HD_MBNT*.docx'))
    p = doan_chua(d, 'Ngày 13 tháng 10')
    thay(p, r'Ngày \d+ tháng \d+ năm \d+', '{{NGAY.GIAI_NGAN_DAY_DU}}', True)
    NH = r'Ngân Hàng TMCP Đầu tư và Phát triển Việt Nam – Chi nhánh Quảng Nam – PGD Hương An'
    for p in d.paragraphs:
        if 'Hợp đồng khung số' in p.text:
            thay(p, '01/2026/MBNT/BIDV-HAK', '{{HOP_DONG.HD_KHUNG_MBNT_SO}}')
            thay(p, '05/01/2026', '{{HOP_DONG.HD_KHUNG_MBNT_NGAY}}')
            thay(p, NH, '{{HOP_DONG.TEN_NGAN_HANG}} – {{HOP_DONG.CHI_NHANH}}', True) if False else thay(p, re.escape(NH), '{{HOP_DONG.TEN_NGAN_HANG}} – {{HOP_DONG.CHI_NHANH}}', True)
            thay(p, r'Công ty TNHH Hoàng Anh Khôi', '{{CONG_TY.TEN}}', True)
    a, b = d.tables[0].rows[0].cells
    thay(doan_chua(a, 'Chi nhánh:'), r'Chi nhánh: Quảng Nam – PGD Hương An', '{{HOP_DONG.CHI_NHANH}}', True)
    thay(doan_chua(a, 'Địa chỉ:'), 'Xã Xuân Phú, thành phố Đà Nẵng', '{{HOP_DONG.DIA_CHI_CHI_NHANH}}')
    thay(doan_chua(a, 'Người đại diện:'), '…………………………….', '{{HOP_DONG.NGUOI_DAI_DIEN_NH}}')
    thay(doan_chua(a, 'Chức vụ/Số CMT:'), r'…+', '{{HOP_DONG.CHUC_VU_NH}}', True)
    thay(doan_chua(a, 'Giấy ủy quyền'), r'993/GUQ-BIDV\.QNA ngày 01/07/2026', '{{HOP_DONG.GIAY_UY_QUYEN_NH}}', True)
    thay(doan_chua(a, 'Điện thoại:'), '0235.3687888', '{{HOP_DONG.DIEN_THOAI_NH}}')
    thay(doan_chua(b, 'BÊN B:'), 'CÔNG TY TNHH HOÀNG ANH KHÔI', '{{CONG_TY.TEN}}')
    thay(doan_chua(b, 'Địa chỉ:'), 'Thôn Thạch Nham Tây, xã Bà Nà, thành phố Đà Nẵng, Việt Nam', '{{CONG_TY.DIA_CHI}}')
    thay(doan_chua(b, 'Người đại diện:'), 'Võ Văn Phụng', '{{CONG_TY.NGUOI_DAI_DIEN}}')
    thay(doan_chua(b, 'Chức vụ/Số CMT:'), 'Giám đốc', '{{CONG_TY.CHUC_VU}}')
    thay(doan_chua(b, 'Điện thoại:'), 'Điện thoại:', 'Điện thoại: {{CONG_TY.DIEN_THOAI}}')
    g1, g2 = d.tables[1].rows[0].cells
    thay(doan_chua(g1, 'Số tiền giao dịch:'), r'19,372\.34 USD', '{{HO_SO.SO_TIEN}} USD', True)
    thay(doan_chua(g1, 'Tỷ giá:'), r'25,810 USD/VND', '{{HO_SO.TY_GIA}} USD/VND', True)
    thay(doan_chua(g1, 'Số tiền thanh'), r'500,000,095', '{{HO_SO.SO_TIEN_QUY_DOI}}', True)
    thay(doan_chua(g1, 'chữ'), r'Năm trăm triệu.*chẵn\.', '{{HO_SO.SO_TIEN_QUY_DOI_BANG_CHU}}.', True)
    thay(doan_chua(g1, 'Ngày giao dịch:'), r'13/10/2026', '{{HO_SO.NGAY_GIAI_NGAN}}', True)
    thay(doan_chua(g1, 'Mục đích sử dụng ngoại tệ:'), r'Thanh toán chi phí của công ty', '{{HO_SO.MUC_DICH}}', True)
    thay(doan_chua(g2, 'Ngày thanh toán:'), r'13/10/2026', '{{HO_SO.NGAY_GIAI_NGAN}}', True)
    c1, c2 = d.tables[2].rows[0].cells
    thay(doan_chua(c1, 'Bên B ủy quyền'), '5620011163', '{{CONG_TY.TAI_KHOAN_USD}}')
    thay(doan_chua(c2, 'VND vào TK số'), '5620727934', '{{CONG_TY.TAI_KHOAN_VND}}')
    thay(doan_chua(c2, 'theo bảng kê'), r'01\.…/202…/8100173/HĐTDngày …+', '{{HO_SO.SO_GIAY_NHAN_NO}} ngày {{HO_SO.NGAY_GIAI_NGAN}}', True)
    d.save(os.path.join(OUT, 'HDMB_USD.docx'))   # hộp kiểm Mua/Bán (content control) được giữ nguyên


# ---------------------------------------------------------------- UNC.xlsx — sửa ở mức XML để giữ logo/hình/ô ghép
NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
M = '{%s}' % NS['m']


def dat_o(root, ref, text):
    c = root.xpath('//m:c[@r="%s"]' % ref, namespaces=NS)
    assert c, 'không có ô ' + ref
    c = c[0]
    for ch in list(c): c.remove(ch)
    if 't' in c.attrib: del c.attrib['t']
    c.set('t', 'inlineStr')
    is_ = etree.SubElement(c, M + 'is'); t = etree.SubElement(is_, M + 't'); t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')


def dat_o_so(root, ref, text_ph):
    """Ô số: chỉ chứa placeholder; engine ghi số (giữ định dạng số của ô). Ở đây lưu dạng chuỗi nội tuyến."""
    dat_o(root, ref, text_ph)


def unc():
    tmp = tempfile.mkdtemp()
    with zipfile.ZipFile(tim('*UNC.xlsx')) as z: z.extractall(tmp)
    p = os.path.join(tmp, 'xl/worksheets/sheet1.xml')
    root = etree.parse(p).getroot()
    dat_o(root, 'J3', '{{CONG_TY.TEN}}')
    dat_o(root, 'I4', '{{CONG_TY.DIA_CHI}}')
    dat_o(root, 'I5', '{{CONG_TY.TAI_KHOAN_VND}}')
    dat_o(root, 'J6', '{{HOP_DONG.TEN_NGAN_HANG}} – {{HOP_DONG.CHI_NHANH}}')
    dat_o_so(root, 'I7', '{{CT.SO_TIEN_QUY_DOI}}')
    dat_o(root, 'I8', '{{CT.SO_TIEN_QUY_DOI_BANG_CHU}} y./.')
    dat_o(root, 'H12', '{{CT.TEN_NGUOI_HUONG}}')
    dat_o(root, 'H16', '{{CT.SO_TAI_KHOAN}}')
    dat_o(root, 'H17', '{{CT.NGAN_HANG}}')
    dat_o(root, 'H18', '{{CT.NOI_DUNG}}')
    dat_o(root, 'X2', '{{HO_SO.NGAY_GIAI_NGAN}}')
    etree.ElementTree(root).write(p, xml_declaration=True, encoding='UTF-8', standalone=True)
    ra = os.path.join(OUT, 'UNC_1_KHOAN.xlsx')
    if os.path.exists(ra): os.remove(ra)
    with zipfile.ZipFile(ra, 'w', zipfile.ZIP_DEFLATED) as z:
        z.write(os.path.join(tmp, '[Content_Types].xml'), '[Content_Types].xml')
        for dp, _, fs in os.walk(tmp):
            for f in fs:
                full = os.path.join(dp, f); rel = os.path.relpath(full, tmp)
                if rel != '[Content_Types].xml': z.write(full, rel)
    shutil.rmtree(tmp)


# ---------------------------------------------------------------- Bảng kê giải ngân (xlsx, 1 hàng mẫu)
def bangke():
    wb = openpyxl.load_workbook(tim('*BANG_KE_GIAI_NGAN*.xlsx'))
    for n in list(wb.sheetnames):
        if n != 'HAK': del wb[n]
    ws = wb['HAK']
    ws['A1'] = 'Tên đơn vị : {{CONG_TY.TEN}}'
    ws['A2'] = 'Địa chỉ : {{CONG_TY.DIA_CHI}}'
    ws['A3'] = 'MST : {{CONG_TY.MA_SO_THUE}}'
    ws['C4'] = '{{HO_SO.NGAY_GIAI_NGAN}}'
    # Xoá 9 hàng dữ liệu mẫu (8..16): chỉ giữ hàng 7 làm HÀNG MẪU; hàng "Tổng cộng" và chữ ký dời lên 9 hàng.
    chieu_cao = {r: ws.row_dimensions[r].height for r in range(1, 30)}
    for m in list(ws.merged_cells.ranges): ws.unmerge_cells(str(m))
    ws.delete_rows(8, 9)
    for r in range(8, 15): ws.row_dimensions[r].height = chieu_cao.get(r + 9)
    for r in range(15, 24): ws.row_dimensions[r].height = None
    ws.merge_cells('A4:B4'); ws.merge_cells('A10:B10'); ws.merge_cells('E10:F10')
    ws['A7'] = '{{CT.STT}}'; ws['B7'] = '{{CT.TEN_NGUOI_HUONG}}'; ws['C7'] = '{{CT.SO_TAI_KHOAN}}'
    ws['D7'] = '{{CT.NGAN_HANG}}'; ws['E7'] = '{{CT.SO_TIEN_QUY_DOI}}'; ws['F7'] = '{{CT.NOI_DUNG}}'
    ws['E8'] = '=SUM(E7:E7)'   # engine tự giãn vùng cộng theo số hàng thụ hưởng
    ws['A7'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws['E7'].alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)
    ws['C4'].alignment = Alignment(horizontal='center')
    for col in 'BCDF': ws[col + '7'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws['E8'].alignment = Alignment(horizontal='right')
    wb.save(os.path.join(OUT, 'BANG_KE_NHIEU_KHOAN.xlsx'))


if __name__ == '__main__':
    hdtd(); hdmb(); unc(); bangke()
    print('Đã tạo mẫu thật tại', os.path.abspath(OUT), sorted(os.listdir(OUT)))
