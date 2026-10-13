# -*- coding: utf-8 -*-
"""Kiểm tra ĐỘC LẬP các file DOCX do engine sinh ra (đọc lại từ file, không dùng dữ liệu nội bộ của engine).
Chạy: python3 kiem_thu/kiem_tra_docx.py <thư_mục_ra>   (sau khi chạy node kiem_thu/chay_thu.js <thư_mục_ra>)"""
import sys, os, json, re, glob, zipfile, subprocess
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

RA = sys.argv[1]
MAU = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'mau_thu')
DIEU_KHOAN = 'Hợp đồng tín dụng cụ thể này được lập thành 03 bản, là một bộ phận không thể tách rời Hợp đồng tín dụng hạn mức.'
loi_tong = 0
bang = []


def tat_ca_van_ban(doc):
    out = []
    def tbl(t):
        for r in t.rows:
            for c in r.cells:
                out.extend(p.text for p in c.paragraphs)
                for n in c.tables: tbl(n)
    out.extend(p.text for p in doc.paragraphs)
    for t in doc.tables: tbl(t)
    for s in doc.sections:
        out.extend(p.text for p in s.header.paragraphs); out.extend(p.text for p in s.footer.paragraphs)
    return out


def so(s):
    return int(s.replace('.', '')) if re.fullmatch(r'\d{1,3}(\.\d{3})*|\d+', s.strip()) else None


def kiem_tra(jf):
    global loi_tong
    e = json.load(open(jf, encoding='utf8')); ten = e['ten']
    f = os.path.join(RA, ten + '.docx'); loi = []
    # 1. gói hợp lệ + XML đúng cú pháp
    with zipfile.ZipFile(f) as z:
        assert z.testzip() is None
        for n in z.namelist():
            if n.endswith('.xml') or n.endswith('.rels'): etree.fromstring(z.read(n))
    doc = Document(f)                                   # python-docx mở được
    van = tat_ca_van_ban(doc); gop = '\n'.join(van)
    # 2. không còn placeholder
    if '{{' in gop or '}}' in gop: loi.append('còn {{ }}')
    if e['engine']['conLai'] or e['engine']['conNgoac']: loi.append('engine báo còn placeholder')
    # 3. tiền: chữ, không cents
    if re.search(r'cents|\d,00\b|\.00\b|\d+\.\d{2}\b', gop): loi.append('có cents/số thập phân')
    if e['loai'] == 'USD':
        if e['tongBangChu'] not in gop: loi.append('thiếu tiền bằng chữ USD')
        if e['tongTien'] == 25000 and 'Hai mươi lăm ngàn đô la Mỹ' not in gop: loi.append('không đúng "Hai mươi lăm ngàn đô la Mỹ"')
    elif e['tongBangChu'] not in gop: loi.append('thiếu tiền bằng chữ VND')
    if e['tongBangSo'] not in gop: loi.append('thiếu tổng bằng số ' + e['tongBangSo'])
    # 4. bảng thụ hưởng
    n = e['soDong']; bt = None
    for t in doc.tables:
        if any('Số tiền' in c.text for c in t.rows[0].cells): bt = t; break
    if bt is None:
        if n != 1: loi.append('không tìm thấy bảng thụ hưởng')
        if 'Người hưởng: ' + e['cacDong'][0]['ten'] not in gop: loi.append('UNC thiếu người hưởng')
    else:
        so_hang = len(bt.rows)
        if so_hang != n + 2: loi.append(f'số hàng bảng {so_hang} != {n}+2')
        tong = 0; thieu_ten = 0
        for i in range(n):
            cells = [c.text for c in bt.rows[1 + i].cells]
            ten_dong = e['cacDong'][i]['ten']
            if not any(ten_dong in c for c in cells): thieu_ten += 1
            tien = [so(c) for c in cells if so(c) is not None]
            # cột số tiền (USD/VND) là số khớp dòng i
            if so(e['cacDong'][i]['soTien']) not in tien: loi.append(f'dòng {i+1} sai số tiền')
            tong += so(e['cacDong'][i]['soTien']) or 0
        if thieu_ten: loi.append(f'{thieu_ten} dòng sai tên')
        if tong != e['tongTien']: loi.append(f'tổng các dòng {tong} != {e["tongTien"]}')
        hang_tong = [c.text for c in bt.rows[-1].cells]
        if e['tongBangSo'] not in hang_tong: loi.append('hàng tổng sai')
        # 5. giữ định dạng: căn phải ở ô tiền mọi hàng; ô ghép hàng tổng; hàng tiêu đề lặp
        mau_doc = Document(os.path.join(MAU, e['mau'] + '.docx'))
        mt = [t for t in mau_doc.tables if any('Số tiền' in c.text for c in t.rows[0].cells)][0]
        cot_tien = [i for i, c in enumerate(mt.rows[1].cells) if mt.rows[1].cells[i].paragraphs[0].alignment == 2][0]
        for i in range(n):
            if bt.rows[1 + i].cells[cot_tien].paragraphs[0].alignment != 2: loi.append(f'dòng {i+1} mất căn phải')
        if bt._tbl.xpath('.//w:tr[1]//w:tblHeader') == []: loi.append('mất tblHeader')
        if not bt.rows[-1]._tr.xpath('.//w:gridSpan'): loi.append('mất ô ghép hàng tổng')
        h_mau = mt.rows[1]._tr.xpath('.//w:trHeight'); h_ra = bt.rows[n]._tr.xpath('.//w:trHeight')
        if h_mau and (not h_ra or h_ra[0].get(qn('w:val')) != h_mau[0].get(qn('w:val'))): loi.append('mất độ cao hàng')
    # 6. nội dung cố định + header/footer + khổ giấy giống mẫu
    if e['mau'].startswith('DE_NGHI') and DIEU_KHOAN not in gop: loi.append('điều khoản cố định bị đổi')
    if e['mau'].startswith('DE_NGHI') and '0126/2026/HDTD-BIDV' not in doc.sections[0].header.paragraphs[0].text: loi.append('header chưa điền')
    if e.get('maHoSo', 'HS0001') not in doc.sections[0].footer.paragraphs[0].text: loi.append('footer chưa điền')
    s_ra, s_mau = doc.sections[0], Document(os.path.join(MAU, e['mau'] + '.docx')).sections[0]
    for a in ('page_width', 'page_height', 'left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        if getattr(s_ra, a) != getattr(s_mau, a): loi.append('khổ giấy/lề đổi: ' + a)
    # 7. PDF qua LibreOffice
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', RA, f], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    pdf = os.path.join(RA, ten + '.pdf'); trang = 0; pdf_ok = os.path.exists(pdf)
    if pdf_ok:
        txt = subprocess.run(['pdftotext', '-layout', pdf, '-'], capture_output=True, text=True).stdout
        trang = txt.count('\f')
        if e['tongBangSo'] not in txt: loi.append('PDF thiếu tổng')
        if '{{' in txt: loi.append('PDF còn {{')
        if bt is not None and trang >= 2 and n >= 5:          # bảng tràn nhiều trang: tiêu đề phải lặp lại ở mỗi trang
            if txt.count('Số tiền') < trang - (0 if n >= 20 else 1): loi.append('hàng tiêu đề không lặp ở các trang sau')
        for i in range(n):
            if e['cacDong'][i]['soTien'] not in txt: loi.append(f'PDF thiếu dòng {i+1}'); break
    else: loi.append('không tạo được PDF')
    loi_tong += len(loi)
    bang.append((ten, n, e['tongBangSo'], trang, 'ĐẠT' if not loi else 'LỖI: ' + '; '.join(loi)))


for jf in sorted(glob.glob(os.path.join(RA, '*.json'))): kiem_tra(jf)
print(f'{"kịch bản":16}{"dòng":>5}  {"tổng":>14}  {"trang PDF":>9}  kết quả')
for b in bang: print(f'{b[0]:16}{b[1]:>5}  {b[2]:>14}  {b[3]:>9}  {b[4]}')
print('TỔNG LỖI:', loi_tong)
sys.exit(1 if loi_tong else 0)
