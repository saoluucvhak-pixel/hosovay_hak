# -*- coding: utf-8 -*-
"""Sinh HUONG_DAN_THIET_LAP.docx (hướng dẫn thiết lập cho người dùng cuối). python3 kiem_thu/tao_huong_dan_docx.py [file_ra]"""
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ra = sys.argv[1] if len(sys.argv) > 1 else 'HUONG_DAN_THIET_LAP.docx'
d = Document()
s = d.sections[0]; s.left_margin = s.right_margin = Cm(2); s.top_margin = s.bottom_margin = Cm(2)
st = d.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
for n in ('Heading 1', 'Heading 2'):
    d.styles[n].font.name = 'Times New Roman'; d.styles[n].font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

_dem = [0]
def h1(t):
    _dem[0] = 0; d.add_heading(t, 1)
def h2(t): d.add_heading(t, 2)
def p(t, b=False):
    x = d.add_paragraph(); r = x.add_run(t); r.bold = b; return x
def li(t, kieu='List Number'):
    if kieu == 'List Number':
        _dem[0] += 1; x = d.add_paragraph(); x.paragraph_format.left_indent = Cm(0.9); x.paragraph_format.first_line_indent = Cm(-0.9)
        r = x.add_run('Bước %d. ' % _dem[0]); r.bold = True; x.add_run(t)
    else:
        d.add_paragraph(t, style=kieu)
def bang(hang, rong=None):
    t = d.add_table(rows=0, cols=len(hang[0])); t.style = 'Table Grid'; t.autofit = False
    for i, h in enumerate(hang):
        c = t.add_row().cells
        for j, v in enumerate(h):
            c[j].text = v
            if i == 0:
                for r in c[j].paragraphs[0].runs: r.bold = True
    if rong:
        for r in t.rows:
            for j, w in enumerate(rong): r.cells[j].width = Cm(w)
    d.add_paragraph()

t = d.add_heading('HƯỚNG DẪN THIẾT LẬP VÀ SỬ DỤNG MẪU VĂN BẢN', 0); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
p('Phần mềm Hồ sơ vay ngân hàng (hosovay_hak) — in chứng từ từ mẫu Word/Excel', True).alignment = WD_ALIGN_PARAGRAPH.CENTER

h1('1. Ý tưởng chung')
p('Bạn sửa biểu mẫu bằng Word/Excel, tải lên phần mềm, phần mềm tự điền số liệu hồ sơ vào bản sao rồi xuất DOCX/XLSX và PDF. '
  'Khi ngân hàng đổi biểu mẫu hoặc có biểu mẫu mới, chỉ cần tải mẫu mới — không phải sửa mã nguồn.')
bang([['Chứng từ', 'Khi nào dùng', 'Định dạng mẫu'],
      ['Văn bản đề nghị giải ngân (HĐTD cụ thể)', 'Mọi hồ sơ; có mẫu riêng cho VND và USD, 1 người và nhiều người thụ hưởng', 'Word (.docx)'],
      ['Ủy nhiệm chi (UNC)', 'Hồ sơ có đúng 1 người thụ hưởng', 'Excel (.xlsx)'],
      ['Bảng kê giải ngân', 'Hồ sơ có từ 2 người thụ hưởng', 'Excel (.xlsx)'],
      ['Hợp đồng mua bán ngoại tệ giao ngay', 'Chỉ hồ sơ USD', 'Word (.docx)']], [5, 8, 3.5])

h1('2. Cài đặt lần đầu (người quản trị làm một lần)')
li('Chép toàn bộ các file .gs và .html của dự án vào Google Apps Script (hoặc dùng clasp). Xoá các file cũ DocxEngine.gs, MauDocx.gs, PhatHanh.gs nếu còn.')
li('Mở Google Sheet → menu "Hồ Sơ Vay NH" → "Khởi tạo / kiểm tra cấu trúc Sheet". Hệ thống tạo các tab DM_MauVanBan, LichSuPhatHanhTaiLieu, NhatKyHeThong và thêm cột liên kết DOCX/PDF/XLSX. Dữ liệu cũ không bị xoá.')
li('Menu "Cấu hình quản trị viên": nhập email những người được tải/đổi mẫu (cách nhau dấu phẩy). Khi chưa cấu hình, ai cũng là quản trị và giao diện hiện cảnh báo đỏ.')
li('Deploy → Manage deployments → New version. Nên đặt quyền truy cập theo tên miền công ty (không để "Anyone"), nếu không phần mềm không nhận ra email người dùng.')
li('Lần chạy đầu tiên Google yêu cầu cấp quyền (Drive, Sheets, UrlFetch) — chọn Cho phép.')

h1('3. Tải mẫu có sẵn')
p('Các mẫu đã chuẩn bị nằm trong thư mục mau_that/ của dự án. Vào Cài đặt → Quản lý mẫu văn bản → khung "Tải mẫu văn bản lên", tải từng file với thông tin sau:')
bang([['File', 'Mã mẫu (gợi ý)', 'Loại chứng từ', 'Loại tiền', 'Điều kiện số người'],
      ['HDTD_USD_1_KHOAN.docx', 'HDTD_USD_1', 'Đề nghị giải ngân', 'USD', '1'],
      ['HDTD_USD_NHIEU_KHOAN.docx', 'HDTD_USD_N', 'Đề nghị giải ngân', 'USD', '2+'],
      ['HDTD_VND_1_KHOAN.docx', 'HDTD_VND_1', 'Đề nghị giải ngân', 'VND', '1'],
      ['HDTD_VND_NHIEU_KHOAN.docx', 'HDTD_VND_N', 'Đề nghị giải ngân', 'VND', '2+'],
      ['UNC_1_KHOAN.xlsx', 'UNC_BIDV', 'Chứng từ thanh toán', 'Tất cả', '1'],
      ['BANG_KE_NHIEU_KHOAN.xlsx', 'BANG_KE_BIDV', 'Chứng từ thanh toán', 'Tất cả', '2+'],
      ['HDMB_USD.docx', 'HDMB_BIDV', 'Hợp đồng mua bán ngoại tệ', 'USD', '1+']], [5.6, 3, 3.4, 1.8, 2.2])
p('Ô "Ngân hàng áp dụng" điền BIDV (hoặc để trống = mọi ngân hàng). Sau khi tải, bấm "Xem trước" ở dòng mẫu để kiểm tra bằng dữ liệu thử — hệ thống tạo file + PDF xem thử, không ảnh hưởng hồ sơ thật.')
p('Điều kiện số người: "1" = đúng 1 người; "2-999" hoặc "2+" = từ 2 người trở lên; "1+" = mọi số lượng.')

h1('4. Thêm mẫu mới từ file Word có sẵn (không cần gõ {{...}})')
p('Dùng khi có biểu mẫu mới của ngân hàng. Làm trong Cài đặt → Quản lý mẫu văn bản → khung "Tạo mẫu từ file Word có sẵn".', False)
li('Chuẩn bị file Word (.docx) biểu mẫu ĐÃ ĐIỀN số liệu thật (hoặc lấy chính chứng từ vừa phát hành). Không cần gõ ký hiệu gì thêm.')
li('Chọn file → bấm "Phân tích file". Phần mềm liệt kê từng đoạn văn kèm vị trí (ví dụ "Bảng 2, hàng 2, ô 5").')
li('Các chỗ nền vàng (✨) là gợi ý tự động: trùng dữ liệu công ty/hợp đồng đang lưu, nhãn thường gặp (Thời hạn vay, rút số tiền vay là, Bằng chữ, ngày … tháng … năm …) và các cột của bảng thụ hưởng. Bấm ✓ để nhận từng gợi ý, hoặc "Áp dụng tất cả gợi ý". Bấm ✕ để bỏ.')
li('Chỗ chưa có gợi ý: dùng chuột bôi chọn đúng đoạn chữ cần thay → gõ/chọn "Trường dữ liệu" (ví dụ CONG_TY.TEN) → bấm "Gán vùng đang chọn". Chỉ đặt con trỏ (không bôi chọn) rồi gán = CHÈN trường tại vị trí đó (ví dụ sau nhãn "Tài liệu liên quan: ").')
li('Bảng thụ hưởng: giữ 1 hàng dữ liệu làm hàng mẫu (hàng này được nhân bản theo số người). Tích "xoá cả hàng bảng này" ở các hàng ví dụ còn lại. KHÔNG xoá hàng tiêu đề và hàng tổng cộng. Ở hàng tổng, gán ô tổng tiền bằng BANG_KE.TONG_TIEN hoặc BANG_KE.TONG_TIEN_QUY_DOI.')
li('Bấm "Tạo mẫu & kiểm tra". Có thể tải file mẫu về mở bằng Word để xem. Nếu hợp lệ bấm "Dùng làm file mẫu ở khung Tải mẫu văn bản".')
li('Kéo lên khung "Tải mẫu văn bản lên": điền Mã mẫu, Tên mẫu, Ngân hàng, Loại tiền, Loại chứng từ, Điều kiện số người → "Tải lên & kiểm tra mẫu". Xong.')
p('Lưu ý: trình tạo mẫu chỉ nhận file .docx. Với file Excel, gõ trực tiếp {{...}} vào ô (xem mục 5). Chữ nằm trong hộp văn bản (text box) không được nhận ra.')

h1('5. Tự gõ placeholder trong Word/Excel (cách thủ công)')
p('Gõ {{NHÓM.TÊN}} (chữ HOA) vào chỗ cần điền. Danh mục đầy đủ có ngay trong phần mềm (Cài đặt → Quản lý mẫu văn bản → Danh mục placeholder). Các nhóm chính:')
bang([['Nhóm', 'Ví dụ', 'Ý nghĩa'],
      ['CONG_TY', '{{CONG_TY.TEN}}, {{CONG_TY.MA_SO_THUE}}, {{CONG_TY.TAI_KHOAN_VND}}', 'Thông tin công ty vay'],
      ['HOP_DONG', '{{HOP_DONG.SO_HOP_DONG}}, {{HOP_DONG.NGAY_KY}}, {{HOP_DONG.CHI_NHANH}}', 'Hợp đồng tín dụng hạn mức, ngân hàng'],
      ['HO_SO', '{{HO_SO.SO_TIEN}}, {{HO_SO.SO_TIEN_BANG_CHU}}, {{HO_SO.THOI_HAN_VAY}}, {{HO_SO.NGAY_GIAI_NGAN}}', 'Thông tin lần giải ngân'],
      ['CT', '{{CT.STT_2}}, {{CT.TEN_NGUOI_HUONG}}, {{CT.SO_TAI_KHOAN}}, {{CT.SO_TIEN_QUY_DOI}}', 'Từng người thụ hưởng — đặt trong 1 hàng bảng để được nhân bản'],
      ['BANG_KE', '{{BANG_KE.TONG_TIEN}}, {{BANG_KE.TONG_TIEN_QUY_DOI}}', 'Tổng cộng'],
      ['NGÀY', '{{NGAY.GIAI_NGAN_DAY_DU}}, {{NGAY_IN}}', 'Ngày tháng']], [2.5, 8, 6])
li('Bảng thụ hưởng: đặt các {{CT.*}} trong MỘT hàng; hàng đó nhân bản theo số người, giữ nguyên font, viền, căn lề, hàng tiêu đề lặp khi sang trang.', 'List Bullet')
li('Excel: ô chỉ chứa 1 placeholder tiền được ghi dạng SỐ (giữ định dạng số, công thức SUM tự giãn theo số hàng, ô ghép/vùng in tự dời).', 'List Bullet')
li('UNC/Bảng kê luôn tính bằng VND quy đổi: dùng {{CT.SO_TIEN_QUY_DOI}} cho cả hồ sơ VND và USD; một mẫu dùng chung cho hai loại tiền.', 'List Bullet')
li('Chữ đỏ đánh dấu trong biểu mẫu gốc được giữ nguyên (là định dạng của mẫu). Muốn in màu đen: đổi chữ chứa trường sang màu Automatic rồi tải lên thành phiên bản mới.', 'List Bullet')

h1('6. Quy tắc tiền và ngày')
li('VND và USD luôn là số nguyên, không có số lẻ/cents. Dấu chấm là ngăn cách hàng nghìn (25.000). Chỉ tỷ giá được có phần thập phân (ví dụ 25.810 hoặc 26.145,5).', 'List Bullet')
li('Số VND quy đổi = làm tròn (USD × tỷ giá bán) về số nguyên; không cộng trực tiếp VND với USD.', 'List Bullet')
li('Tiền canh phải, ngày canh giữa, chữ canh trái; ngày dạng dd/mm/yyyy.', 'List Bullet')
li('Hồ sơ cũ còn số lẻ: vẫn xem được, sửa thông tin khác không đụng số tiền, sửa số tiền phải nhập số nguyên; phát hành chứng từ bị chặn và nêu rõ mã hồ sơ, dòng, trường, giá trị cần sửa.', 'List Bullet')

h1('7. Phát hành chứng từ')
li('Mở hồ sơ → tab "Thực hiện" → khu "Danh sách chứng từ hồ sơ".')
li('Mỗi dòng: Tên tài liệu, Mã mẫu, Phiên bản, Ngân hàng, Loại tiền, Trạng thái, Ngày tạo, Người tạo. Nút: Xem trước · Tạo chứng từ/Tạo lại · DOCX · PDF · XLSX · Google · Lịch sử · Thư mục.')
li('"Tạo trọn bộ" sinh đủ chứng từ theo quy tắc: 1 người → UNC; nhiều người → Bảng kê; USD → thêm Hợp đồng mua bán ngoại tệ; VND → không có hợp đồng này.')
li('Mỗi lần tạo sinh file mới (_L1, _L2…); bản cũ chuyển "Đã thay thế" và không bị xoá/ghi đè. Bấm đúp trong 15 giây không tạo trùng.')
li('Phần mềm chặn phát hành khi thiếu dữ liệu bắt buộc và nêu rõ {{KHOÁ}} cùng nơi cần bổ sung (ví dụ DM_CongTy.NguoiDaiDien).')
li('Tệp lưu trong Drive: HoSoVay / <Mã hồ sơ> / DOCX | XLSX | PDF; lịch sử ghi ở tab LichSuPhatHanhTaiLieu (kể cả lần lỗi).')

h1('8. Quản lý mẫu hằng ngày')
bang([['Thao tác', 'Tác dụng'],
      ['Xem trước', 'Điền dữ liệu thử (hoặc mã hồ sơ thật nếu nhập) và tạo file + PDF xem thử; không phát hành'],
      ['Kiểm tra', 'Kiểm tra placeholder và file mẫu còn nguyên vẹn (hash). Nếu ai sửa file mẫu ngoài hệ thống, phát hành bị chặn'],
      ['Ngừng / Khôi phục', 'Ngừng dùng một phiên bản; khôi phục phiên bản cũ khi cần'],
      ['Mặc định', 'Ưu tiên mẫu này khi nhiều mẫu cùng phù hợp'],
      ['Tải lại cùng Mã mẫu', 'Tạo phiên bản mới; bản cũ chuyển "Hết hiệu lực", vẫn giữ file'],
      ['Lịch sử', 'Xem ai tải/đổi mẫu, lúc nào']], [4.5, 12])
p('Cách chọn mẫu khi phát hành: loại chứng từ + loại tiền (cụ thể hơn "Tất cả") + ngân hàng (cụ thể hơn để trống) + mẫu mặc định + khoảng số người hẹp hơn + phiên bản cao hơn.')

h1('9. Khi gặp lỗi')
bang([['Thông báo / hiện tượng', 'Cách xử lý'],
      ['Thiếu {{CONG_TY.…}} / {{HOP_DONG.…}}', 'Bổ sung dữ liệu ở danh mục công ty/hợp đồng theo nguồn thông báo, rồi tạo lại'],
      ['Mẫu cho nhiều người nhưng không có hàng {{CT.*}}', 'Đặt các {{CT.*}} vào một hàng bảng, hoặc đổi điều kiện = 1'],
      ['Placeholder không có trong danh mục', 'Kiểm tra chính tả khoá (chữ HOA, dấu chấm) theo danh mục trong phần mềm'],
      ['File mẫu đã bị thay đổi (hash khác)', 'Tải lại mẫu thành phiên bản mới'],
      ['Hồ sơ có số lẻ', 'Sửa số tiền về số nguyên rồi phát hành'],
      ['Không có mẫu phù hợp', 'Hệ thống dùng biểu mẫu mặc định dựng bằng code; thêm mẫu theo mục 4 để dùng mẫu của bạn']], [6.5, 10])

h1('10. Quay lại phiên bản cũ')
li('Mã nguồn: Deploy → Manage deployments → chọn phiên bản trước → Deploy. Dữ liệu không cần khôi phục (tab/cột mới chỉ thêm).')
li('Quay về biểu mẫu dựng bằng code mà không đổi mã: Ngừng toàn bộ mẫu trong Quản lý mẫu văn bản.')
li('Chứng từ đã phát hành không bị ảnh hưởng.')

h1('11. Lưu ý trung thực')
p('Phần mềm đã được kiểm thử bằng môi trường giả lập Google, LibreOffice và trình duyệt Chromium. Chưa chạy trên Google Apps Script thật và chưa mở bằng Microsoft Word/Excel thật. '
  'Trước khi dùng thật, hãy tạo một hồ sơ nháp, bấm "Tạo trọn bộ", mở các file DOCX/XLSX/PDF bằng Word/Excel và so với biểu mẫu gốc.')
d.save(ra); print('Đã tạo', ra)
